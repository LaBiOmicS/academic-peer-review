#!/usr/bin/env python3
# /// script
# requires-python = ">=3.10"
# dependencies = [
#     "pypdf>=4.0.0",
#     "pdfplumber>=0.11.0",
#     "python-docx>=1.1.0",
# ]
# ///
"""Multi-format academic paper extractor, structural analyzer, 3-Axis Classifier & Security Sanitizer.

Classifies manuscripts across:
- Axis 1: 5 Major Scientific Areas (`--area`)
- Axis 2: 15 Academic Article Types (`--article-type`)
- Axis 3: 12 Research Approaches / Designs (`--research-approach`)

Information Security Hardening:
- Validates regular file type and maximum file size (`--max-file-mb`, default 100 MB).
- Uses `--` option-termination and `timeout=60` on external subprocesses (`pdftotext`, `pandoc --sandbox`).
  Pandoc is only ever run with `--sandbox` (pandoc >= 2.15); there is no unsandboxed fallback.
- Strips invisible/bidirectional Unicode control characters and scans for embedded prompt-injection patterns.
"""

from __future__ import annotations

import argparse
import json
import re
import shutil
import subprocess
import sys
from pathlib import Path
from typing import Any

sys.path.insert(0, str(Path(__file__).resolve().parent))
from area_profiles import (  # noqa: E402
    ARTICLE_TYPE_DEFINITIONS,
    RESEARCH_APPROACH_DEFINITIONS,
    classify_manuscript,
    sanitize_and_scan_manuscript,
)
from journal_scope import clean_journal_name  # noqa: E402


SECTION_PATTERNS = {
    "abstract": re.compile(r"(?im)^\s*(?:#+\s*|\d+\.?\s*)?(?:abstract|resumo|resumen)\b"),
    "introduction": re.compile(r"(?im)^\s*(?:#+\s*|[1I]+\.?\s*)?(?:introduction|background|introdução|referencial\s+teórico)\b"),
    "methods": re.compile(
        r"(?im)^\s*(?:#+\s*|\d+\.?\s*)?(?:methods|methodology|materials\s+and\s+methods|experimental\s+setup|approach|search\s+strategy|data\s+sources|metodologia|materiais\s+e\s+métodos|procedimentos\s+metodológicos)\b"
    ),
    "results": re.compile(
        r"(?im)^\s*(?:#+\s*|\d+\.?\s*)?(?:results|experiments|empirical\s+evaluation|findings|data\s+records|technical\s+validation|case\s+presentation|resultados|análise\s+e\s+discussão|relato\s+do\s+caso)\b"
    ),
    "discussion": re.compile(
        r"(?im)^\s*(?:#+\s*|\d+\.?\s*)?(?:discussion|conclusion|conclusions|usage\s+notes|discussão|considerações\s+finais|conclusão|conclusões)\b"
    ),
    "limitations": re.compile(
        r"(?im)^\s*(?:#+\s*|\d+\.?\s*)?(?:limitations|broader\s+impact|threats\s+to\s+validity|limitações)\b"
    ),
    "reproducibility": re.compile(
        r"(?im)(?:data\s+availability|code\s+availability|reproducibility\s+statement|open\s+science|disponibilidade\s+de\s+dados|availability\s+and\s+requirements)"
    ),
    "ethics": re.compile(
        r"(?im)(?:ethics\s+statement|ethical\s+approval|institutional\s+review\s+board|\bIRB\b|competing\s+interests|conflict\s+of\s+interest|comitê\s+de\s+ética|\bCEP\b|\bCONEP\b|\bCAAE\b|\bCEUA\b)"
    ),
    "references": re.compile(
        r"(?im)^\s*(?:#+\s*|\d+\.?\s*)?(?:references|bibliography|literature\s+cited|referências|referências\s+bibliográficas)\b"
    ),
}

DOI_REGEX = re.compile(r"\b(10\.\d{4,9}/[-._;()/:A-Za-z0-9]+)\b")
# The "arXiv:" / arxiv.org prefix is required: bare numbers such as "1203.4567" are usually not IDs.
ARXIV_REGEX = re.compile(
    r"(?:\barXiv:\s*|arxiv\.org/(?:abs|pdf)/)((?:0[7-9]|[12]\d)(?:0[1-9]|1[0-2])\.\d{4,5}(?:v\d+)?)\b",
    re.IGNORECASE,
)
PMID_REGEX = re.compile(r"\bPMID:\s*(\d{6,9})\b", re.IGNORECASE)
ACCESSION_REGEX = re.compile(
    r"\b(GSE\d{3,8}|GSM\d{3,8}|PRJNA\d{4,9}|PRJEB\d{4,9}|SRR\d{5,10}|ERR\d{5,10}|PXD\d{5,9}|NCT\d{8}|RBR-[a-z0-9]{5,10}|CRD\d{8,12})\b"
)
FIGURE_CAPTION_REGEX = re.compile(r"(?im)^\s*(?:Figure|Fig\.?|Figura)\s*(\d+[A-Za-z]?)[\s.:—-]+([^\n]{10,250})")
TABLE_CAPTION_REGEX = re.compile(r"(?im)^\s*(?:Table|Tab\.?|Tabela|Quadro)\s*(\d+[A-Za-z]?)[\s.:—-]+([^\n]{10,250})")


def _is_white(color: Any) -> bool:
    """True for white fill in gray (1), RGB (1,1,1) or CMYK (0,0,0,0) colour spaces."""
    if color is None:
        return False
    values = tuple(color) if isinstance(color, (list, tuple)) else (color,)
    try:
        values = tuple(float(v) for v in values)
    except (TypeError, ValueError):
        return False
    if len(values) in (1, 3):
        return all(v >= 0.98 for v in values)
    if len(values) == 4:
        return all(v <= 0.02 for v in values)
    return False


def _find_hidden_pdf_text(path: Path, max_pages: int = 60, limit: int = 15) -> list[str]:
    """Return runs of characters drawn white or at <= 2pt (invisible to a human reader).

    White-on-white text is the most common carrier of prompt injections aimed at AI reviewers.
    Requires pdfplumber; returns [] when it is unavailable.
    """
    try:
        import pdfplumber  # type: ignore
    except ImportError:
        return []

    snippets: list[str] = []
    try:
        with pdfplumber.open(str(path)) as pdf:
            for page in pdf.pages[:max_pages]:
                current: list[str] = []
                for ch in page.chars:
                    hidden = _is_white(ch.get("non_stroking_color")) or float(ch.get("size") or 99) <= 2.0
                    if hidden:
                        current.append(ch.get("text", ""))
                    elif current:
                        snippets.append("".join(current))
                        current = []
                if current:
                    snippets.append("".join(current))
    except Exception:  # noqa: BLE001 - detection is best effort and must never block extraction
        return []

    cleaned = [" ".join(s.split())[:120] for s in snippets]
    return [s for s in cleaned if len(s) >= 4][:limit]


def extract_pdf(path: Path) -> tuple[str, dict[str, Any]]:
    meta: dict[str, Any] = {"backend": None, "pages": 0, "hidden_text_runs": _find_hidden_pdf_text(path)}

    try:
        import fitz  # type: ignore

        doc = fitz.open(str(path))
        meta["backend"] = "pymupdf"
        meta["pages"] = len(doc)
        pages_text = []
        for idx, page in enumerate(doc, start=1):
            pages_text.append(f"\n--- [Page {idx}] ---\n" + page.get_text("text"))
        text = "\n".join(pages_text)
        if len(text.strip()) >= 500:
            return text, meta
    except Exception:
        pass

    try:
        import pdfplumber  # type: ignore

        with pdfplumber.open(str(path)) as pdf:
            meta["backend"] = "pdfplumber"
            meta["pages"] = len(pdf.pages)
            pages_text = []
            for idx, page in enumerate(pdf.pages, start=1):
                pages_text.append(f"\n--- [Page {idx}] ---\n" + (page.extract_text() or ""))
            text = "\n".join(pages_text)
            if len(text.strip()) >= 500:
                return text, meta
    except Exception:
        pass

    try:
        from pypdf import PdfReader  # type: ignore

        reader = PdfReader(str(path))
        meta["backend"] = "pypdf"
        meta["pages"] = len(reader.pages)
        pages_text = []
        for idx, page in enumerate(reader.pages, start=1):
            pages_text.append(f"\n--- [Page {idx}] ---\n" + (page.extract_text() or ""))
        text = "\n".join(pages_text)
        if len(text.strip()) >= 500:
            return text, meta
    except Exception:
        pass

    if shutil.which("pdftotext"):
        try:
            proc = subprocess.run(
                ["pdftotext", "-layout", "--", str(path), "-"],
                capture_output=True,
                text=True,
                check=True,
                timeout=60,
            )
            meta["backend"] = "pdftotext"
            meta["pages"] = proc.stdout.count("\f") + 1
            return proc.stdout, meta
        except Exception:
            pass

    raise RuntimeError(
        f"Failed to extract text from PDF {path}. Ensure pypdf, pymupdf, or pdftotext is installed."
    )


def _find_hidden_docx_runs(doc: Any, limit: int = 15) -> list[str]:
    """Return text of runs a human reader would not see (hidden, white, or ~1pt font).

    Invisible text is the classic carrier for prompt injection aimed at AI reviewers.
    """
    from docx.shared import Pt, RGBColor  # type: ignore

    paragraphs = list(doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.extend(cell.paragraphs)

    hidden: list[str] = []
    for para in paragraphs:
        for run in para.runs:
            text = run.text.strip()
            if not text:
                continue
            font = run.font
            try:
                is_white = font.color is not None and font.color.type is not None and font.color.rgb == RGBColor(0xFF, 0xFF, 0xFF)
            except (AttributeError, ValueError):
                is_white = False
            is_tiny = font.size is not None and font.size <= Pt(2)
            if font.hidden or is_white or is_tiny:
                hidden.append(" ".join(text.split())[:120])
                if len(hidden) >= limit:
                    return hidden
    return hidden


def extract_docx(path: Path) -> tuple[str, dict[str, Any]]:
    meta: dict[str, Any] = {"backend": None}
    try:
        import docx  # type: ignore

        doc = docx.Document(str(path))
        meta["backend"] = "python-docx"
        meta["hidden_text_runs"] = _find_hidden_docx_runs(doc)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        table_texts = []
        for t_idx, table in enumerate(doc.tables, start=1):
            rows = []
            for row in table.rows:
                rows.append(" | ".join(cell.text.strip() for cell in row.cells))
            if rows:
                table_texts.append(f"\n[Table {t_idx}]\n" + "\n".join(rows))
        return "\n\n".join(paragraphs + table_texts), meta
    except Exception:
        pass

    if shutil.which("pandoc"):
        try:
            proc = subprocess.run(
                ["pandoc", "--sandbox", "-t", "markdown", "--", str(path)],
                capture_output=True,
                text=True,
                check=True,
                timeout=60,
            )
            meta["backend"] = "pandoc"
            return proc.stdout, meta
        except Exception:
            pass

    raise RuntimeError(f"Failed to extract DOCX {path}. Install python-docx or pandoc >= 2.15 (needs --sandbox).")


PLAIN_TEXT_FORMATS = {".txt", ".md", ".markdown", ".tex", ".rst"}
PANDOC_FORMATS = {".odt", ".rtf", ".html", ".htm", ".epub"}


def extract_with_pandoc(path: Path) -> tuple[str, dict[str, Any]]:
    if not shutil.which("pandoc"):
        raise RuntimeError(f"Reading {path.suffix} files requires pandoc >= 2.15 (with --sandbox).")
    try:
        proc = subprocess.run(
            ["pandoc", "--sandbox", "-t", "plain", "--wrap=none", "--", str(path)],
            capture_output=True,
            text=True,
            check=True,
            timeout=60,
        )
    except (subprocess.CalledProcessError, subprocess.TimeoutExpired, OSError) as e:
        raise RuntimeError(f"pandoc failed to convert {path.name}: {e}") from e
    return proc.stdout, {"backend": "pandoc"}


def analyze_structure(
    text: str,
    source_path: Path,
    backend_meta: dict[str, Any],
    *,
    security_scan: dict[str, Any],
    domain_hint: str = "",
    explicit_area: str = "auto",
    explicit_article_type: str = "auto",
    explicit_research_approach: str = "auto",
    target_journal: str = "",
) -> dict[str, Any]:
    lines = text.splitlines()
    words = re.findall(r"\S+", text)

    detected_sections = {
        name: bool(pattern.search(text))
        for name, pattern in SECTION_PATTERNS.items()
    }

    dois = sorted({d.rstrip(".,;)]}") for d in DOI_REGEX.findall(text)})
    arxiv_ids = sorted(set(ARXIV_REGEX.findall(text)))
    pmids = sorted(set(PMID_REGEX.findall(text)))
    accessions = sorted(set(ACCESSION_REGEX.findall(text)))

    figures = [
        {"id": m.group(1), "caption": " ".join(m.group(2).split())}
        for m in FIGURE_CAPTION_REGEX.finditer(text)
    ][:40]
    tables = [
        {"id": m.group(1), "caption": " ".join(m.group(2).split())}
        for m in TABLE_CAPTION_REGEX.finditer(text)
    ][:40]

    non_empty = [
        ln.strip()
        for ln in lines[:40]
        if ln.strip() and not ln.strip().startswith("--- [Page")
    ]
    inferred_title = non_empty[0] if non_empty else source_path.stem

    # The inferred name comes from untrusted manuscript text: it is reduced to inert
    # journal-title characters and only accepted from an explicit "submitted to:"-style line.
    inferred_journal = clean_journal_name(target_journal)
    journal_source = "user" if inferred_journal else None
    if not inferred_journal:
        j_match = re.search(
            r"(?im)^\s*(?:submitted\s+to|target\s+journal|intended\s+for|submetido\s+(?:a|à|ao)|revista\s+alvo)\s*[:—\-]\s*([^\n\r]{3,80})$",
            text[:3000],
        )
        if j_match:
            inferred_journal = clean_journal_name(j_match.group(1))
            journal_source = "manuscript" if inferred_journal else None

    is_likely_academic = (
        len(words) >= 400
        and sum(1 for k in ("abstract", "introduction", "methods", "results", "discussion", "references") if detected_sections.get(k)) >= 2
    )

    classification = classify_manuscript(
        text,
        domain_hint=domain_hint,
        explicit_area=explicit_area,
        explicit_article_type=explicit_article_type,
        explicit_research_approach=explicit_research_approach,
    )

    return {
        "source_file": str(source_path.resolve()),
        "format": source_path.suffix.lower(),
        "inferred_title": inferred_title[:300],
        "target_journal": inferred_journal or None,
        "target_journal_source": journal_source,
        "char_count": len(text),
        "word_count": len(words),
        "line_count": len(lines),
        **backend_meta,
        "is_likely_academic_paper": is_likely_academic,
        "security_scan": security_scan,
        "scientific_area": classification,
        "detected_sections": detected_sections,
        "figures_detected": len(figures),
        "tables_detected": len(tables),
        "figure_captions": figures,
        "table_captions": tables,
        "identifiers": {
            "dois": dois,
            "arxiv_ids": arxiv_ids,
            "pmids": pmids,
            "accessions_and_registries": accessions,
        },
    }


def main() -> int:
    parser = argparse.ArgumentParser(description="Extract text, structural metadata, security scan, and 3-axis classification.")
    parser.add_argument("input_path", help="Path to PDF, DOCX, LaTeX (.tex), Markdown (.md), or text file.")
    parser.add_argument("--output-text", required=True, help="Where to write the extracted plain text.")
    parser.add_argument("--output-meta", help="Optional path to write structural diagnostic JSON.")
    parser.add_argument(
        "--area",
        choices=["auto", "biologicas", "saude", "exatas", "engenharias", "humanas"],
        default="auto",
        help="Major Scientific Area override or auto-detect (default: auto).",
    )
    parser.add_argument(
        "--article-type",
        choices=["auto"] + list(ARTICLE_TYPE_DEFINITIONS.keys()),
        default="auto",
        help="Academic Article Type override or auto-detect (default: auto).",
    )
    parser.add_argument(
        "--research-approach",
        choices=["auto"] + list(RESEARCH_APPROACH_DEFINITIONS.keys()),
        default="auto",
        help="Research Approach override or auto-detect (default: auto).",
    )
    parser.add_argument("--domain", default="", help="Optional specific domain hint.")
    parser.add_argument("--journal", "--target-journal", dest="journal", default="", help="Optional target journal name.")
    parser.add_argument("--max-file-mb", type=float, default=100.0, help="Maximum allowed input file size in MB (default: 100).")
    args = parser.parse_args()

    src = Path(args.input_path).resolve()
    if not src.exists() or not src.is_file():
        print(f"Error: Input path is not a regular file: {src}", file=sys.stderr)
        return 1

    size_mb = src.stat().st_size / (1024 * 1024)
    if size_mb > args.max_file_mb:
        print(
            f"Error: Input file size ({size_mb:.1f} MB) exceeds maximum allowed limit ({args.max_file_mb} MB).",
            file=sys.stderr,
        )
        return 1

    ext = src.suffix.lower()
    try:
        if ext == ".pdf":
            raw_text, backend_meta = extract_pdf(src)
        elif ext == ".docx":
            raw_text, backend_meta = extract_docx(src)
        elif ext in PANDOC_FORMATS:
            raw_text, backend_meta = extract_with_pandoc(src)
        elif ext in PLAIN_TEXT_FORMATS:
            raw_text = src.read_text(encoding="utf-8", errors="replace")
            backend_meta = {"backend": "direct_read", "pages": max(1, len(raw_text) // 3000)}
        else:
            print(
                f"Error: unsupported input format '{ext or src.name}'. Supported: .pdf, .docx, "
                f"{', '.join(sorted(PANDOC_FORMATS | PLAIN_TEXT_FORMATS))}. Convert legacy .doc files to .docx or PDF first.",
                file=sys.stderr,
            )
            return 1
    except RuntimeError as e:
        print(f"Error: {e}", file=sys.stderr)
        return 1

    if "\x00" in raw_text[:65536]:
        print(f"Error: '{src.name}' looks like a binary file, not text. Convert it to PDF, DOCX or plain text.", file=sys.stderr)
        return 1

    text, security_scan = sanitize_and_scan_manuscript(raw_text)
    hidden_runs = backend_meta.pop("hidden_text_runs", None) or []
    if hidden_runs:
        security_scan["hidden_formatted_text_runs"] = len(hidden_runs)
        security_scan["suspicious_patterns"] = [
            {"category": "hidden_formatted_text", "matched_text": t} for t in hidden_runs
        ] + security_scan["suspicious_patterns"]
        security_scan["suspicious_prompt_injection_count"] += len(hidden_runs)
        security_scan["is_clean"] = False

    out_text = Path(args.output_text).resolve()
    out_text.parent.mkdir(parents=True, exist_ok=True)
    out_text.write_text(text, encoding="utf-8")

    structure = analyze_structure(
        text,
        src,
        backend_meta,
        security_scan=security_scan,
        domain_hint=args.domain,
        explicit_area=args.area,
        explicit_article_type=args.article_type,
        explicit_research_approach=args.research_approach,
        target_journal=args.journal,
    )
    if args.output_meta:
        out_meta = Path(args.output_meta).resolve()
        out_meta.parent.mkdir(parents=True, exist_ok=True)
        out_meta.write_text(json.dumps(structure, indent=2, ensure_ascii=False), encoding="utf-8")

    print(json.dumps({
        "status": "ok",
        "char_count": structure["char_count"],
        "word_count": structure["word_count"],
        "pages": structure.get("pages", 0),
        "backend": structure.get("backend"),
        "is_likely_academic_paper": structure["is_likely_academic_paper"],
        "target_journal": structure["target_journal"],
        "target_journal_source": structure["target_journal_source"],
        "security_scan": security_scan,
        "classification": {
            "area_code": structure["scientific_area"]["area_code"],
            "area_name": structure["scientific_area"]["area_name"],
            "article_type": structure["scientific_area"]["article_type"]["type_code"],
            "research_approach": structure["scientific_area"]["research_approach"]["approach_code"],
            "transversal_flags": structure["scientific_area"]["transversal_flags"],
        },
        "detected_sections": [k for k, v in structure["detected_sections"].items() if v],
        "figures_detected": structure["figures_detected"],
        "tables_detected": structure["tables_detected"],
        "dois_found": len(structure["identifiers"]["dois"]),
        "arxiv_ids_found": len(structure["identifiers"]["arxiv_ids"]),
        "accessions_found": len(structure["identifiers"]["accessions_and_registries"]),
    }, indent=2, ensure_ascii=False))

    if structure["char_count"] < 1000:
        print(
            "WARNING: Extracted text is shorter than 1,000 characters. The PDF may be scanned/image-only. "
            "Use Gemini's native multimodal read_file tool directly on the PDF if OCR text is missing.",
            file=sys.stderr,
        )
        return 2

    return 0


if __name__ == "__main__":
    sys.exit(main())
